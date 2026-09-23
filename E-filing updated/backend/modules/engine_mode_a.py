from docx import Document
from docx.shared import Cm, Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

def generate_mode_a_document(data: dict) -> io.BytesIO:
    doc = Document()
    
    # Global margins: Left 5cm, Right 3cm, Top 1.5 inches, Bottom 1.5 inches
    for section in doc.sections:
        section.left_margin = Cm(5)
        section.right_margin = Cm(3)
        section.top_margin = Inches(1.5)
        section.bottom_margin = Inches(1.5)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Stamp Paper Logic
    is_affidavit = data.get("is_affidavit_or_deed", False)
    
    # 1. Cause Title Block
    title = doc.add_paragraph()
    if is_affidavit:
        title.paragraph_format.space_before = Inches(4.5) - Inches(1.5) # Total 4.5 inches from top
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(f"IN THE {data.get('court_name', 'HON\'BLE COURT').upper()}\n")
    title_run.bold = True
    
    case_num = doc.add_paragraph(f"CASE NO. {data.get('case_number', '__________')} OF {data.get('year', '20__')}")
    case_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("\n")
    
    party_a = doc.add_paragraph(f"{data.get('petitioner_name', 'Petitioner Name')} ... Petitioner")
    vs = doc.add_paragraph("VERSUS")
    vs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    party_b = doc.add_paragraph(f"{data.get('respondent_name', 'Respondent Name')} ... Respondent")

    doc.add_paragraph("\n")
    
    # 2. Synopsis & List of Dates
    if data.get('synopsis'):
        doc.add_heading("SYNOPSIS & LIST OF DATES", level=1)
        for act in data.get('synopsis', []):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{act.get('date', 'Date')}: ").bold = True
            p.add_run(f"{act.get('event', 'Event Description')}")
        doc.add_page_break()

    # 3. Body Paragraphs (Auto-numbered)
    doc.add_heading("PETITION UNDER RELEVANT PROVISIONS OF LAW", level=1)
    paragraphs = data.get('body_paragraphs', ["The petitioner respectfully submits as follows:"])
    for idx, para_text in enumerate(paragraphs, 1):
        p = doc.add_paragraph(f"{idx}. {para_text}")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 4. The Prayer Clause
    doc.add_paragraph("\n")
    prayer = doc.add_paragraph("Wherefore it is most respectfully prayed that this Hon'ble Court may be pleased to:")
    
    for clause in data.get('prayer_clauses', ["Pass such other orders as deemed fit."]):
        p = doc.add_paragraph(clause, style='List Number')
        
    doc.add_paragraph("\n")
    footer = doc.add_paragraph(f"Place: {data.get('place', '_________')}\nDate: {data.get('date', '_________')}")
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    sign = doc.add_paragraph("ADVOCATE FOR PETITIONER")
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
