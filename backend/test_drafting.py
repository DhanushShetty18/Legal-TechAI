"""Tests for the Camera Filing drafting engine.

Run from the ``backend`` directory:  python -m pytest test_drafting.py -v
"""

import io
import json
import re
import zipfile
from unittest.mock import MagicMock, patch

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from modules.drafting import (
    DRAFTABLE_CASE_TYPES,
    SALE_DEED,
    SALE_DEED_FIELD_KEYS,
    build_sections,
    document_by_id,
    estimate_pages,
    extract_from_document,
    merge_extractions,
    missing_required_fields,
    render_plain_text,
    requirements_for,
)
from modules.drafting import checklists, extraction, formatting, stream
from modules.drafting.docx_export import render_sale_deed_docx
from modules.drafting.pdf import render_sale_deed_pdf
from modules.drafting.schema import DERIVABLE_FIELDS, SALE_DEED_FIELDS
from routers.drafting import router as drafting_router

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

FULL_FIELDS = {
    "place_of_execution": "Mangaluru",
    "date_of_execution": "2026-03-15",
    "vendor_name": "Ramesh Kumar",
    "vendor_father_name": "Suresh Kumar",
    "vendor_age": "52",
    "vendor_marital_status": "Married",
    "vendor_occupation": "Business",
    "vendor_nationality": "Indian",
    "vendor_address": "No. 14, Kadri Road, Mangaluru 575002, Karnataka",
    "vendor_pan": "ABCDE1234F",
    "vendor_aadhaar": "123412341234",
    "purchaser_name": "Anita Rao",
    "purchaser_father_name": "Vasudev Rao",
    "purchaser_age": "38",
    "purchaser_marital_status": "Married",
    "purchaser_occupation": "Medical Practitioner",
    "purchaser_nationality": "Indian",
    "purchaser_address": "Flat 3B, Lalbagh Residency, Mangaluru 575003",
    "purchaser_pan": "ZYXWV9876K",
    "purchaser_aadhaar": "987698769876",
    "prior_deed_document_no": "MNG-1-04521/2009",
    "prior_deed_book_no": "Book-I",
    "prior_deed_volume_no": "184",
    "prior_deed_pages": "112 to 126",
    "prior_deed_date": "2009-07-21",
    "prior_deed_sro": "SR Mangaluru",
    "sale_consideration_amount": "4512000",
    "market_value": "4600000",
    "stamp_duty_paid": "235000",
    "registration_fee": "45120",
    "mode_of_payment": "RTGS dated 10.03.2026",
    "property_description": "A residential site with the dwelling house thereon.",
    "survey_number": "128/3",
    "sub_division_number": "2",
    "pt_sheet_number": "PT-44/9",
    "property_area": "2,400 sq. ft.",
    "village_or_city": "Mangaluru",
    "taluka": "Mangaluru",
    "district": "Dakshina Kannada",
    "state": "Karnataka",
    "local_authority": "Mangaluru City Corporation",
    "boundary_east": "Site No. 129",
    "boundary_west": "Site No. 127",
    "boundary_north": "30 ft. wide road",
    "boundary_south": "Property of Sri Krishna Bhat",
    "witness_1_name": "Ganesh Shenoy",
    "witness_1_address": "Bejai, Mangaluru",
    "witness_2_name": "Latha Nayak",
    "witness_2_address": "Kankanady, Mangaluru",
}


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(drafting_router, prefix="/drafting")
    return TestClient(app)


def gemini_returning(payload):
    """A patch target for genai.GenerativeModel that answers with ``payload``."""
    model = MagicMock()
    response = MagicMock()
    response.text = payload if isinstance(payload, str) else json.dumps(payload)
    model.generate_content.return_value = response
    factory = MagicMock(return_value=model)
    return factory, model


# ---------------------------------------------------------------------------
# formatting
# ---------------------------------------------------------------------------

class TestFormatting:
    @pytest.mark.parametrize("value,expected", [
        (0, "Zero"),
        (7, "Seven"),
        (15, "Fifteen"),
        (40, "Forty"),
        (52, "Fifty Two"),
        (100, "One Hundred"),
        (345, "Three Hundred Forty Five"),
        (1000, "One Thousand"),
        (4512000, "Forty Five Lakh Twelve Thousand"),
        (10000000, "One Crore"),
        (112500000, "Eleven Crore Twenty Five Lakh"),
    ])
    def test_number_to_words_uses_indian_grouping(self, value, expected):
        assert formatting.number_to_words(value) == expected

    def test_rupees_in_words_is_deed_phrasing(self):
        assert formatting.rupees_in_words(4512000) == (
            "Rupees Forty Five Lakh Twelve Thousand Only"
        )

    def test_rupees_in_words_handles_paise(self):
        assert "Fifty Paise" in formatting.rupees_in_words(1000.50)

    def test_rupees_in_words_of_nothing_is_empty(self):
        assert formatting.rupees_in_words(None) == ""

    @pytest.mark.parametrize("value,expected", [
        (100, "100"),
        (1000, "1,000"),
        (45120, "45,120"),
        (4512000, "45,12,000"),
        (112500000, "11,25,00,000"),
    ])
    def test_format_indian_currency(self, value, expected):
        assert formatting.format_indian_currency(value) == expected

    @pytest.mark.parametrize("raw,expected", [
        ("4500000", 4500000.0),
        ("Rs. 45,00,000/-", 4500000.0),
        ("₹ 45,00,000", 4500000.0),
        (4500000, 4500000.0),
        ("", None),
        (None, None),
        ("not a number", None),
    ])
    def test_parse_amount(self, raw, expected):
        assert formatting.parse_amount(raw) == expected

    @pytest.mark.parametrize("raw,expected", [
        ("2026-03-15", ("15th", "March", "2026")),
        ("01-04-2026", ("1st", "April", "2026")),
        ("22/11/2025", ("22nd", "November", "2025")),
        ("03.12.2024", ("3rd", "December", "2024")),
    ])
    def test_deed_date_parts(self, raw, expected):
        assert formatting.deed_date_parts(raw) == expected

    def test_unparseable_date_becomes_fill_in_lines(self):
        day, month, year = formatting.deed_date_parts("sometime last spring")
        assert set(day) == {"_"} and set(month) == {"_"}
        assert year == "20____"

    @pytest.mark.parametrize("day,suffix", [
        (11, "11th"), (12, "12th"), (13, "13th"), (21, "21st"), (23, "23rd"),
    ])
    def test_ordinals_handle_the_teens(self, day, suffix):
        assert formatting.deed_date_parts(f"2026-01-{day:02d}")[0] == suffix

    def test_aadhaar_is_masked_to_last_four_digits(self):
        assert formatting.mask_aadhaar("1234 5678 9012") == "XXXX XXXX 9012"
        assert formatting.mask_aadhaar("123456789012") == "XXXX XXXX 9012"

    def test_mask_aadhaar_of_junk_is_empty(self):
        assert formatting.mask_aadhaar("n/a") == ""
        assert formatting.mask_aadhaar(None) == ""

    def test_normalise_pan(self):
        assert formatting.normalise_pan("abcde1234f") == "ABCDE1234F"
        assert formatting.normalise_pan("ABCDE 1234 F") == "ABCDE1234F"
        assert formatting.normalise_pan("garbage") == "garbage"


# ---------------------------------------------------------------------------
# checklists and schema
# ---------------------------------------------------------------------------

class TestChecklists:
    def test_sale_deed_is_offered_and_draftable(self):
        assert SALE_DEED in checklists.REQUIREMENTS_BY_CASE_TYPE
        assert SALE_DEED in DRAFTABLE_CASE_TYPES

    def test_divorce_case_is_still_offered(self):
        assert checklists.DIVORCE_CASE in checklists.REQUIREMENTS_BY_CASE_TYPE
        assert requirements_for(checklists.DIVORCE_CASE)

    def test_sale_deed_checklist_covers_every_phase_of_conveyancing(self):
        phases = requirements_for(SALE_DEED)
        assert len(phases) == 5
        joined = " ".join(p["phase"] for p in phases).lower()
        for topic in ("identity", "title", "clearance", "stamp duty", "witness"):
            assert topic in joined

    def test_document_ids_are_unique(self):
        ids = [d["id"] for p in requirements_for(SALE_DEED) for d in p["documents"]]
        assert len(ids) == len(set(ids))

    def test_every_extracts_entry_names_a_real_field(self):
        for phase in requirements_for(SALE_DEED):
            for doc in phase["documents"]:
                unknown = set(doc["extracts"]) - set(SALE_DEED_FIELD_KEYS)
                assert not unknown, f"{doc['id']} names unknown field(s) {unknown}"

    def test_every_required_field_has_a_document_that_supplies_it(self):
        """The flow must not require a particular no upload can produce."""
        supplied = {
            key
            for phase in requirements_for(SALE_DEED)
            for doc in phase["documents"]
            for key in doc["extracts"]
        }
        for field in SALE_DEED_FIELDS:
            if field.required and field.key not in DERIVABLE_FIELDS:
                # Execution place and date are asked of the user directly.
                if field.key in ("place_of_execution", "date_of_execution"):
                    continue
                assert field.key in supplied, f"nothing supplies {field.key}"

    def test_unknown_case_type_has_no_requirements(self):
        assert requirements_for("Maritime Salvage") == []

    def test_document_by_id(self):
        assert document_by_id(SALE_DEED, "parent_sale_deed")["required"] is True
        assert document_by_id(SALE_DEED, "nope") == {}

    def test_field_definitions_are_serialisable_for_the_ui(self):
        payload = [f.as_dict() for f in SALE_DEED_FIELDS]
        assert len(payload) == len(SALE_DEED_FIELD_KEYS)
        assert all({"key", "label", "group", "required", "kind"} <= set(f)
                   for f in payload)
        json.dumps(payload)


# ---------------------------------------------------------------------------
# extraction
# ---------------------------------------------------------------------------

class TestExtraction:
    def test_prompt_is_scoped_to_the_documents_own_fields(self):
        prompt = extraction.build_prompt("Vendor Aadhaar", ["vendor_name", "vendor_pan"])
        assert "vendor_name" in prompt and "vendor_pan" in prompt
        assert "boundary_east" not in prompt
        assert "Vendor Aadhaar" in prompt

    def test_extract_from_document_keeps_only_known_fields(self):
        factory, _ = gemini_returning({
            "vendor_name": "Ramesh Kumar",
            "vendor_pan": "ABCDE1234F",
            "favourite_colour": "blue",
            "raw_text": "PAN ABCDE1234F  Ramesh Kumar",
        })
        with patch("google.generativeai.GenerativeModel", factory):
            result = extract_from_document(b"jpeg", "image/jpeg", "vendor_id")
        assert result["fields"]["vendor_name"] == "Ramesh Kumar"
        assert "favourite_colour" not in result["fields"]
        assert result["error"] is None

    @pytest.mark.parametrize("empty", ["", "  ", "null", "N/A", "not found", "-", None])
    def test_placeholder_answers_are_dropped(self, empty):
        factory, _ = gemini_returning({"vendor_name": empty, "raw_text": ""})
        with patch("google.generativeai.GenerativeModel", factory):
            result = extract_from_document(b"jpeg", "image/jpeg", "vendor_id")
        assert "vendor_name" not in result["fields"]

    def test_aadhaar_is_masked_at_the_point_of_extraction(self):
        factory, _ = gemini_returning({
            "vendor_aadhaar": "1234 5678 9012",
            "raw_text": "Aadhaar 1234 5678 9012",
        })
        with patch("google.generativeai.GenerativeModel", factory):
            result = extract_from_document(b"jpeg", "image/jpeg", "vendor_id")
        assert result["fields"]["vendor_aadhaar"] == "XXXX XXXX 9012"
        assert "5678" not in result["fields"]["vendor_aadhaar"]

    def test_mistranscribed_pan_is_rescued_from_the_raw_text(self):
        fields = extraction.reconcile_identifiers(
            {"vendor_pan": "ABCD1234F"},           # only four leading letters
            "PERMANENT ACCOUNT NUMBER ABCDE1234F",
        )
        assert fields["vendor_pan"] == "ABCDE1234F"

    def test_a_valid_pan_answer_is_left_alone(self):
        fields = extraction.reconcile_identifiers(
            {"vendor_pan": "ABCDE1234F"}, "ABCDE1234F and ZZZZZ9999Z"
        )
        assert fields["vendor_pan"] == "ABCDE1234F"

    def test_pan_is_recovered_when_the_model_missed_it_entirely(self):
        fields = extraction.reconcile_identifiers({}, "PAN: ABCDE1234F")
        assert fields["vendor_pan"] == "ABCDE1234F"

    def test_a_vendors_pan_is_never_attributed_to_the_purchaser(self):
        """A vendor's ID card carries one PAN; it belongs to the vendor alone.

        Attributing it to both parties puts a wrong identifier for a party into
        a registered deed, and it passes the review form as already filled.
        """
        fields = extraction.reconcile_identifiers(
            {},
            "PERMANENT ACCOUNT NUMBER ABCDE1234F",
            allowed_keys=document_by_id(SALE_DEED, "vendor_id")["extracts"],
        )
        assert fields["vendor_pan"] == "ABCDE1234F"
        assert "purchaser_pan" not in fields

    def test_a_purchasers_document_fills_only_the_purchasers_identifiers(self):
        fields = extraction.reconcile_identifiers(
            {},
            "PAN ABCDE1234F  AADHAAR 1234 5678 9012",
            allowed_keys=document_by_id(SALE_DEED, "purchaser_id")["extracts"],
        )
        assert fields["purchaser_pan"] == "ABCDE1234F"
        assert fields["purchaser_aadhaar"] == "XXXX XXXX 9012"
        assert "vendor_pan" not in fields
        assert "vendor_aadhaar" not in fields

    def test_a_document_about_neither_party_gains_no_identifiers(self):
        fields = extraction.reconcile_identifiers(
            {},
            "PAN ABCDE1234F",
            allowed_keys=document_by_id(SALE_DEED, "witness_ids")["extracts"],
        )
        assert "vendor_pan" not in fields
        assert "purchaser_pan" not in fields

    def test_extraction_scopes_identifiers_to_the_uploaded_documents_party(self):
        """The end-to-end path, which is where this was actually caught."""
        factory, _ = gemini_returning({
            "vendor_name": "Ramesh Kumar",
            "raw_text": "Ramesh Kumar  PAN ABCDE1234F",
        })
        with patch("google.generativeai.GenerativeModel", factory):
            result = extract_from_document(b"jpeg", "image/jpeg", "vendor_id")
        assert result["fields"]["vendor_pan"] == "ABCDE1234F"
        assert "purchaser_pan" not in result["fields"]

    def test_extraction_failure_is_reported_not_raised(self):
        factory = MagicMock(side_effect=RuntimeError("quota exhausted"))
        with patch("google.generativeai.GenerativeModel", factory), \
             patch("modules.drafting.extraction.time.sleep" if hasattr(extraction, "time")
                   else "time.sleep", MagicMock()):
            result = extract_from_document(b"jpeg", "image/jpeg", "vendor_id")
        assert result["fields"] == {}
        assert "quota exhausted" in result["error"]

    def test_malformed_json_is_reported_not_raised(self):
        factory, _ = gemini_returning("this is not json at all")
        with patch("google.generativeai.GenerativeModel", factory), \
             patch("time.sleep", MagicMock()):
            result = extract_from_document(b"jpeg", "image/jpeg", "vendor_id")
        assert result["error"] is not None
        assert result["fields"] == {}

    def test_authoritative_document_wins_over_an_incidental_mention(self):
        merged = merge_extractions([
            {"docId": "property_tax_receipt",
             "fields": {"vendor_address": "Old address off a tax receipt"}},
            {"docId": "vendor_address_proof",
             "fields": {"vendor_address": "Current address off the utility bill"}},
        ])
        assert merged["fields"]["vendor_address"].startswith("Current address")
        assert merged["provenance"]["vendor_address"] == "vendor_address_proof"

    def test_first_non_empty_wins_among_equals(self):
        merged = merge_extractions([
            {"docId": "utility_bills", "fields": {"property_description": "First"}},
            {"docId": "property_tax_receipt", "fields": {"property_description": "Second"}},
        ])
        assert merged["fields"]["property_description"] == "First"

    def test_amounts_are_normalised_to_plain_digits(self):
        merged = merge_extractions([
            {"docId": "payment_proof",
             "fields": {"sale_consideration_amount": "Rs. 45,12,000/-"}},
        ])
        assert merged["fields"]["sale_consideration_amount"] == "4512000"

    def test_missing_required_fields_is_empty_for_a_complete_set(self):
        assert missing_required_fields(FULL_FIELDS) == []

    def test_missing_required_fields_lists_what_the_user_must_supply(self):
        missing = missing_required_fields({})
        assert "vendor_name" in missing
        assert "boundary_east" in missing

    def test_derivable_fields_are_never_asked_for(self):
        missing = missing_required_fields({})
        assert not DERIVABLE_FIELDS & set(missing)

    def test_extraction_summary_reports_readiness(self):
        assert extraction.extraction_summary(FULL_FIELDS)["readyToGenerate"] is True
        summary = extraction.extraction_summary({})
        assert summary["readyToGenerate"] is False
        assert summary["fieldsFilled"] == 0
        assert summary["fieldsTotal"] == len(SALE_DEED_FIELD_KEYS)


# ---------------------------------------------------------------------------
# generator
# ---------------------------------------------------------------------------

class TestGenerator:
    def test_sections_follow_the_model_draft_in_order(self):
        ids = [s["id"] for s in build_sections(FULL_FIELDS)]
        expected_order = [
            "title", "preamble", "between_label", "vendor", "and_label",
            "purchaser", "expression", "title_recital", "agreement_recital",
            "witnesseth", "clause_1",
        ]
        assert ids[:len(expected_order)] == expected_order
        # The closing sequence the Sub-Registrar's office looks for.
        assert ids[-6:] == [
            "schedule_body", "boundaries", "attestation", "sign_vendor",
            "sign_purchaser", "witnesses",
        ]

    def test_clauses_are_numbered_contiguously_from_one(self):
        numbers = [int(s["heading"]) for s in build_sections(FULL_FIELDS)
                   if s["style"] == "clause"]
        assert numbers == list(range(1, len(numbers) + 1))
        assert len(numbers) >= 13, "a registrable deed carries at least 13 clauses"

    def test_every_style_is_one_the_renderers_know(self):
        known = {
            "title", "preamble", "party-label", "party", "recital",
            "operative-heading", "clause", "schedule-heading", "schedule-note",
            "schedule", "boundaries", "attestation", "signature-block",
            "witnesses",
        }
        assert {s["style"] for s in build_sections(FULL_FIELDS)} <= known

    def test_the_deed_carries_the_particulars_it_was_given(self):
        text = render_plain_text(build_sections(FULL_FIELDS))
        assert "RAMESH KUMAR" in text
        assert "ANITA RAO" in text
        assert "MNG-1-04521/2009" in text
        assert "Dakshina Kannada" in text
        assert "Ganesh Shenoy" in text

    def test_consideration_appears_in_figures_and_in_words(self):
        text = render_plain_text(build_sections(FULL_FIELDS))
        assert "Rs. 45,12,000/-" in text
        assert "Rupees Forty Five Lakh Twelve Thousand Only" in text

    def test_consideration_words_supplied_by_the_user_are_respected(self):
        fields = dict(FULL_FIELDS,
                      sale_consideration_words="Rupees Forty Five Lakh Twelve Thousand Only (agreed)")
        assert "(agreed)" in render_plain_text(build_sections(fields))

    def test_the_deed_never_reproduces_a_full_aadhaar_number(self):
        text = render_plain_text(build_sections(FULL_FIELDS))
        assert "123412341234" not in text.replace(" ", "")
        assert "XXXX XXXX 1234" in text
        assert not re.search(r"\b\d{4}\s?\d{4}\s?\d{4}\b", text)

    def test_the_statutory_clauses_carry_the_reference_deeds_language(self):
        text = render_plain_text(build_sections(FULL_FIELDS))
        for phrase in (
            "NOW THIS DEED WITNESSETH AS UNDER:-",
            "NO OBJECTION CERTIFICATE",
            "free from all sorts of encumbrances",
            "IN WITNESS WHEREOF",
            "SCHEDULE",
            "they are Indian",
        ):
            assert phrase in text, f"missing {phrase!r}"

    def test_an_empty_field_set_still_produces_a_complete_deed(self):
        sections = build_sections({})
        assert len(sections) == len(build_sections(FULL_FIELDS))
        text = render_plain_text(sections)
        assert "SALE DEED" in text
        assert "_____" in text, "unknown particulars must appear as fill-in lines"

    def test_partial_particulars_do_not_disturb_the_structure(self):
        sections = build_sections({"vendor_name": "Ramesh Kumar"})
        assert [s["id"] for s in sections] == [s["id"] for s in build_sections({})]
        assert "RAMESH KUMAR" in render_plain_text(sections)

    def test_optional_property_identifiers_are_omitted_when_absent(self):
        without = build_sections({k: v for k, v in FULL_FIELDS.items()
                                  if k != "pt_sheet_number"})
        schedule = next(s for s in without if s["id"] == "schedule_body")
        assert "P.T. Sheet No." not in schedule["body"]

    def test_narrative_override_replaces_only_the_bespoke_passages(self):
        narrative = {"title_recital": "WHEREAS this recital came from the model."}
        sections = {s["id"]: s["body"] for s in build_sections(FULL_FIELDS, narrative)}
        assert sections["title_recital"] == narrative["title_recital"]
        assert "doth hereby sell" in sections["clause_2"], "clauses stay locked"

    def test_estimated_length_is_in_the_five_to_ten_page_band(self):
        assert 5 <= estimate_pages(build_sections(FULL_FIELDS)) <= 10

    def test_render_plain_text_numbers_the_clauses(self):
        text = render_plain_text(build_sections(FULL_FIELDS))
        assert "1. That in pursuance" in text


# ---------------------------------------------------------------------------
# streaming
# ---------------------------------------------------------------------------

class TestStreaming:
    def test_deltas_reassemble_the_body_exactly(self):
        body = "WHEREAS  the Vendor\nis the owner.\n\nEast: Site No. 129\n"
        assert "".join(stream.iter_deltas(body)) == body

    def test_deltas_of_an_empty_body_are_empty(self):
        assert list(stream.iter_deltas("")) == []

    def test_every_section_arrives_and_every_delta_belongs_to_it(self):
        events = list(stream.stream_document(FULL_FIELDS, use_llm=False))
        assembled = {}
        for event in events:
            if event["type"] == "delta":
                assembled[event["id"]] = assembled.get(event["id"], "") + event["text"]
        for event in events:
            if event["type"] == "section_end":
                assert assembled.get(event["id"], "") == event["body"]

    def test_the_stream_opens_with_metadata_and_closes_with_the_whole_deed(self):
        events = list(stream.stream_document(FULL_FIELDS, use_llm=False))
        types = [e["type"] for e in events]
        assert types.index("meta") < types.index("section")
        assert types[-1] == "done"
        done = events[-1]
        assert len(done["sections"]) == len(build_sections(FULL_FIELDS))
        assert "SALE DEED" in done["plainText"]
        assert 5 <= done["estimatedPages"] <= 10

    def test_sections_stream_in_document_order(self):
        events = list(stream.stream_document(FULL_FIELDS, use_llm=False))
        streamed = [e["id"] for e in events if e["type"] == "section"]
        assert streamed == [s["id"] for s in build_sections(FULL_FIELDS)]

    def test_every_event_is_json_serialisable(self):
        for event in stream.stream_document(FULL_FIELDS, use_llm=False):
            json.dumps(event)

    def test_missing_particulars_raise_a_warning_but_not_a_failure(self):
        events = list(stream.stream_document({"vendor_name": "Ramesh"}, use_llm=False))
        warnings = [e for e in events if e["type"] == "warning"]
        assert warnings and warnings[0]["missingRequired"]
        assert events[-1]["type"] == "done"

    def test_a_complete_set_raises_no_warning(self):
        events = list(stream.stream_document(FULL_FIELDS, use_llm=False))
        assert not [e for e in events if e["type"] == "warning"]

    def test_model_authored_passages_are_streamed_when_available(self):
        narrative = {
            "title_recital": "WHEREAS " + "the Vendor is the absolute owner " * 6,
            "agreement_recital": "AND WHEREAS " + "the Vendor has agreed to sell " * 6,
            "schedule_body": "Survey No.: 128/3\n\n" + "All that piece and parcel " * 6,
        }
        with patch.object(stream, "author_narrative", return_value=narrative):
            events = list(stream.stream_document(FULL_FIELDS, use_llm=True))
        done = events[-1]
        assert done["source"] == "gemini"
        bodies = {s["id"]: s["body"] for s in done["sections"]}
        assert bodies["title_recital"] == narrative["title_recital"]
        assert bodies["schedule_body"] == narrative["schedule_body"]

    def test_a_failed_model_call_falls_back_to_the_precedent_wording(self):
        with patch.object(stream, "author_narrative",
                          side_effect=RuntimeError("503 model overloaded")):
            events = list(stream.stream_document(FULL_FIELDS, use_llm=True))
        done = events[-1]
        assert done["source"] == "template"
        assert done["type"] == "done"
        bodies = {s["id"]: s["body"] for s in done["sections"]}
        assert bodies["title_recital"].startswith("WHEREAS")
        assert "MNG-1-04521/2009" in bodies["title_recital"]
        assert [e for e in events if e["type"] == "notice"]

    def test_a_failed_model_call_still_yields_the_full_deed(self):
        with patch.object(stream, "author_narrative", side_effect=RuntimeError("boom")):
            events = list(stream.stream_document(FULL_FIELDS, use_llm=True))
        assert len(events[-1]["sections"]) == len(build_sections(FULL_FIELDS))

    def test_author_narrative_rejects_stub_passages(self):
        factory, _ = gemini_returning({
            "title_recital": "See attached.",
            "agreement_recital": "As agreed.",
            "schedule_body": "The property.",
        })
        with patch("google.generativeai.GenerativeModel", factory), \
             patch("time.sleep", MagicMock()):
            with pytest.raises(Exception):
                stream.author_narrative(FULL_FIELDS)

    def test_author_narrative_keeps_the_passages_that_are_usable(self):
        good = "WHEREAS the Vendor is the absolute owner of the Schedule property " \
               "and acquired the same under a registered Sale Deed of 2009 duly " \
               "registered before the Sub-Registrar."
        factory, _ = gemini_returning({
            "title_recital": good,
            "agreement_recital": "short",
            "schedule_body": "",
        })
        with patch("google.generativeai.GenerativeModel", factory):
            narrative = stream.author_narrative(FULL_FIELDS)
        assert narrative == {"title_recital": good}

    def test_author_narrative_strips_markdown(self):
        good = ("**WHEREAS** the Vendor is the absolute owner of the Schedule "
                "property and acquired the same under a registered Sale Deed "
                "duly registered before the Sub-Registrar of Mangaluru.")
        factory, _ = gemini_returning({"title_recital": good})
        with patch("google.generativeai.GenerativeModel", factory):
            narrative = stream.author_narrative(FULL_FIELDS)
        assert "**" not in narrative["title_recital"]

    def test_the_narrative_prompt_forbids_inventing_particulars(self):
        assert "Never invent" in stream._NARRATIVE_SYSTEM_PROMPT
        assert "operative clauses" in stream._NARRATIVE_SYSTEM_PROMPT

    def test_the_narrative_prompt_carries_the_supplied_particulars(self):
        prompt = stream._narrative_prompt(FULL_FIELDS)
        assert "MNG-1-04521/2009" in prompt
        assert "title_recital" in prompt

    def test_llm_availability_follows_the_api_key(self, monkeypatch):
        monkeypatch.delenv("GEMINI_API_KEY", raising=False)
        assert stream.llm_available() is False
        monkeypatch.setenv("GEMINI_API_KEY", "test-key")
        assert stream.llm_available() is True


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def read_pdf_text(buffer: io.BytesIO):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(buffer.getvalue()))
    return reader, "\n".join(page.extract_text() or "" for page in reader.pages)


class TestPdf:
    def test_a_complete_deed_renders_to_five_to_ten_pages(self):
        buffer = render_sale_deed_pdf(fields=FULL_FIELDS)
        reader, _ = read_pdf_text(buffer)
        assert 5 <= len(reader.pages) <= 10

    def test_the_output_is_a_valid_pdf(self):
        assert render_sale_deed_pdf(fields=FULL_FIELDS).getvalue()[:5] == b"%PDF-"

    def test_the_pdf_carries_the_deeds_language_and_particulars(self):
        _, text = read_pdf_text(render_sale_deed_pdf(fields=FULL_FIELDS))
        for phrase in ("SALE DEED", "NOW THIS DEED WITNESSETH", "RAMESH KUMAR",
                       "ANITA RAO", "SCHEDULE", "Rs. 45,12,000/-",
                       "Dakshina Kannada", "Ganesh Shenoy"):
            assert phrase in text, f"missing {phrase!r}"

    def test_the_pdf_carries_the_impression_and_witness_pages(self):
        _, text = read_pdf_text(render_sale_deed_pdf(fields=FULL_FIELDS))
        assert "SIGNATURE OF THE VENDOR" in text
        assert "SIGNATURE OF THE PURCHASER" in text
        assert "Photograph" in text
        assert "Thumb" in text
        assert "WITNESSES" in text

    def test_pages_are_numbered(self):
        _, text = read_pdf_text(render_sale_deed_pdf(fields=FULL_FIELDS))
        assert "Page 1" in text and "Page 2" in text

    def test_the_approved_sections_are_what_get_printed(self):
        """A user edit on screen must reach the PDF verbatim."""
        sections = [dict(s) for s in build_sections(FULL_FIELDS)]
        for section in sections:
            if section["id"] == "clause_1":
                section["body"] = "That this clause was amended by the advocate."
        _, text = read_pdf_text(render_sale_deed_pdf(sections=sections))
        assert "amended by the advocate" in text

    def test_margins_leave_room_for_binding_on_legal_paper(self):
        from reportlab.lib.units import inch

        from modules.drafting import pdf as pdf_module

        assert pdf_module.LEFT_MARGIN >= 1.5 * inch
        assert pdf_module.RIGHT_MARGIN >= 0.75 * inch
        assert pdf_module.TOP_MARGIN >= 0.75 * inch
        assert pdf_module.BOTTOM_MARGIN >= 0.75 * inch

    def test_stamp_paper_offset_pushes_the_first_page_down(self):
        from reportlab.lib.units import inch

        plain = render_sale_deed_pdf(fields=FULL_FIELDS)
        offset = render_sale_deed_pdf(fields=FULL_FIELDS,
                                      first_page_top_offset=3.5 * inch)
        assert len(read_pdf_text(offset)[0].pages) >= len(read_pdf_text(plain)[0].pages)
        assert "SALE DEED" in read_pdf_text(offset)[1]

    def test_an_empty_field_set_still_renders(self):
        reader, text = read_pdf_text(render_sale_deed_pdf(fields={}))
        assert len(reader.pages) >= 5
        assert "SALE DEED" in text

    def test_markup_in_a_field_cannot_break_the_render(self):
        fields = dict(FULL_FIELDS, vendor_name="Ramesh <b>Kumar</b> & Sons")
        _, text = read_pdf_text(render_sale_deed_pdf(fields=fields))
        assert "<b>" in text or "&lt;b&gt;" not in text
        assert "SALE DEED" in text


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

class TestDocx:
    def test_the_output_is_a_readable_word_document(self):
        buffer = render_sale_deed_docx(fields=FULL_FIELDS)
        assert zipfile.is_zipfile(io.BytesIO(buffer.getvalue()))

        from docx import Document

        document = Document(io.BytesIO(buffer.getvalue()))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "SALE DEED" in text
        assert "NOW THIS DEED WITNESSETH AS UNDER:-" in text
        assert "RAMESH KUMAR" in text
        assert "Rupees Forty Five Lakh Twelve Thousand Only" in text

    def test_the_docx_carries_the_impression_grids(self):
        from docx import Document

        buffer = render_sale_deed_docx(fields=FULL_FIELDS)
        document = Document(io.BytesIO(buffer.getvalue()))
        assert len(document.tables) == 2
        assert "Thumb of Right Hand" in document.tables[0].cell(0, 0).text

    def test_an_empty_field_set_still_renders(self):
        assert render_sale_deed_docx(fields={}).getvalue()


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

class TestApi:
    def test_case_types_reports_what_can_be_drafted(self, client):
        body = client.get("/drafting/case-types").json()
        names = {c["name"]: c for c in body["caseTypes"]}
        assert names[SALE_DEED]["draftable"] is True
        assert names[SALE_DEED]["documentCount"] > 0
        assert "llmAvailable" in body

    def test_requirements_returns_the_checklist_and_the_field_definitions(self, client):
        response = client.get("/drafting/requirements", params={"case_type": SALE_DEED})
        assert response.status_code == 200
        body = response.json()
        assert len(body["phases"]) == 5
        assert len(body["fields"]) == len(SALE_DEED_FIELD_KEYS)
        assert body["draftable"] is True

    def test_requirements_for_the_divorce_case_still_work(self, client):
        body = client.get("/drafting/requirements",
                          params={"case_type": "Divorce Case"}).json()
        assert body["phases"]
        assert body["draftable"] is False

    def test_an_unknown_case_type_is_a_404(self, client):
        assert client.get("/drafting/requirements",
                          params={"case_type": "Maritime Salvage"}).status_code == 404

    def test_preview_returns_the_whole_deed_without_calling_the_model(self, client):
        body = client.post("/drafting/preview", json={"fields": FULL_FIELDS}).json()
        assert len(body["sections"]) == len(build_sections(FULL_FIELDS))
        assert 5 <= body["estimatedPages"] <= 10
        assert body["readyToGenerate"] is True

    def test_preview_of_nothing_reports_what_is_missing(self, client):
        body = client.post("/drafting/preview", json={"fields": {}}).json()
        assert body["readyToGenerate"] is False
        assert "vendor_name" in body["missingRequired"]

    def test_extract_merges_several_uploads_into_one_field_set(self, client):
        factory, model = gemini_returning({
            "vendor_name": "Ramesh Kumar",
            "vendor_pan": "ABCDE1234F",
            "raw_text": "Ramesh Kumar ABCDE1234F",
        })
        with patch("google.generativeai.GenerativeModel", factory):
            response = client.post(
                "/drafting/extract",
                data={"doc_ids": "vendor_id,purchaser_id", "case_type": SALE_DEED},
                files=[
                    ("files", ("a.jpg", b"fake-jpeg-bytes", "image/jpeg")),
                    ("files", ("b.jpg", b"fake-jpeg-bytes", "image/jpeg")),
                ],
            )
        assert response.status_code == 200
        body = response.json()
        assert len(body["documents"]) == 2
        assert body["fields"]["vendor_name"] == "Ramesh Kumar"
        assert body["readyToGenerate"] is False
        assert "vendor_address" in body["missingRequired"]
        assert model.generate_content.call_count == 2

    def test_extract_accepts_a_pdf_upload(self, client):
        factory, _ = gemini_returning({"vendor_name": "Ramesh", "raw_text": ""})
        with patch("google.generativeai.GenerativeModel", factory):
            response = client.post(
                "/drafting/extract",
                data={"doc_ids": "parent_sale_deed"},
                files=[("files", ("deed.pdf", b"%PDF-1.4 fake", "application/pdf"))],
            )
        assert response.status_code == 200

    def test_extract_rejects_a_file_that_is_not_a_document(self, client):
        response = client.post(
            "/drafting/extract",
            data={"doc_ids": "vendor_id"},
            files=[("files", ("notes.txt", b"hello", "text/plain"))],
        )
        assert response.status_code == 400

    def test_extract_rejects_an_empty_file(self, client):
        response = client.post(
            "/drafting/extract",
            data={"doc_ids": "vendor_id"},
            files=[("files", ("a.jpg", b"", "image/jpeg"))],
        )
        assert response.status_code == 400

    def test_extract_rejects_an_oversized_file(self, client):
        from routers import drafting as drafting_module

        oversized = b"x" * (drafting_module.MAX_UPLOAD_BYTES + 1)
        response = client.post(
            "/drafting/extract",
            data={"doc_ids": "vendor_id"},
            files=[("files", ("a.jpg", oversized, "image/jpeg"))],
        )
        assert response.status_code == 413

    def test_extract_rejects_too_many_files_at_once(self, client):
        from routers import drafting as drafting_module

        files = [("files", (f"{i}.jpg", b"bytes", "image/jpeg"))
                 for i in range(drafting_module.MAX_UPLOADS_PER_REQUEST + 1)]
        assert client.post("/drafting/extract", data={"doc_ids": ""},
                           files=files).status_code == 400

    def test_one_unreadable_scan_does_not_sink_the_filing(self, client):
        factory = MagicMock(side_effect=RuntimeError("image too blurred"))
        with patch("google.generativeai.GenerativeModel", factory), \
             patch("time.sleep", MagicMock()):
            response = client.post(
                "/drafting/extract",
                data={"doc_ids": "vendor_id"},
                files=[("files", ("a.jpg", b"bytes", "image/jpeg"))],
            )
        assert response.status_code == 200
        assert response.json()["documents"][0]["error"]

    def test_merge_folds_per_document_results_together(self, client):
        body = client.post("/drafting/merge", json={
            "documents": [
                {"docId": "vendor_id", "fields": {"vendor_name": "Ramesh Kumar"}},
                {"docId": "purchaser_id", "fields": {"purchaser_name": "Anita Rao"}},
            ],
        }).json()
        assert body["fields"] == {"vendor_name": "Ramesh Kumar",
                                  "purchaser_name": "Anita Rao"}
        assert body["provenance"]["vendor_name"] == "vendor_id"
        assert body["readyToGenerate"] is False

    def test_merge_applies_the_documents_precedence_rules(self, client):
        body = client.post("/drafting/merge", json={
            "documents": [
                {"docId": "property_tax_receipt",
                 "fields": {"vendor_address": "Off the tax receipt"}},
                {"docId": "vendor_address_proof",
                 "fields": {"vendor_address": "Off the utility bill"}},
            ],
        }).json()
        assert body["fields"]["vendor_address"] == "Off the utility bill"

    def test_a_user_correction_outranks_every_scan(self, client):
        body = client.post("/drafting/merge", json={
            "documents": [
                {"docId": "vendor_id", "fields": {"vendor_name": "RAMESH KUMER"}},
            ],
            "overrides": {"vendor_name": "Ramesh Kumar"},
        }).json()
        assert body["fields"]["vendor_name"] == "Ramesh Kumar"
        assert body["provenance"]["vendor_name"] == "user"

    def test_an_empty_override_does_not_erase_an_extracted_value(self, client):
        body = client.post("/drafting/merge", json={
            "documents": [
                {"docId": "vendor_id", "fields": {"vendor_name": "Ramesh Kumar"}},
            ],
            "overrides": {"vendor_name": "   "},
        }).json()
        assert body["fields"]["vendor_name"] == "Ramesh Kumar"

    def test_merge_of_nothing_is_not_an_error(self, client):
        body = client.post("/drafting/merge", json={"documents": []}).json()
        assert body["fields"] == {}
        assert body["readyToGenerate"] is False

    def test_merge_reports_readiness_for_a_complete_set(self, client):
        body = client.post("/drafting/merge", json={
            "documents": [], "overrides": FULL_FIELDS,
        }).json()
        assert body["readyToGenerate"] is True

    def _sse_events(self, response):
        events = []
        for line in response.text.splitlines():
            if line.startswith("data: "):
                events.append(json.loads(line[6:]))
        return events

    def test_generate_streams_the_deed_as_server_sent_events(self, client):
        response = client.post("/drafting/generate",
                               json={"fields": FULL_FIELDS, "useLlm": False})
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/event-stream")
        assert response.headers["x-accel-buffering"] == "no"

        events = self._sse_events(response)
        types = [e["type"] for e in events]
        assert "meta" in types and "section" in types and "delta" in types
        assert types[-1] == "done"
        assert "SALE DEED" in events[-1]["plainText"]

    def test_generate_streams_every_section(self, client):
        response = client.post("/drafting/generate",
                               json={"fields": FULL_FIELDS, "useLlm": False})
        streamed = [e["id"] for e in self._sse_events(response)
                    if e["type"] == "section"]
        assert streamed == [s["id"] for s in build_sections(FULL_FIELDS)]

    def test_generate_reports_a_failure_on_the_stream(self, client):
        with patch("routers.drafting.stream_document",
                   side_effect=RuntimeError("engine offline")):
            response = client.post("/drafting/generate", json={"fields": {}})
        events = self._sse_events(response)
        assert events[-1]["type"] == "error"
        assert "engine offline" in events[-1]["message"]

    def test_pdf_export_returns_a_downloadable_pdf(self, client):
        response = client.post("/drafting/export/pdf",
                               json={"fields": FULL_FIELDS,
                                     "filename": "Sale Deed Ramesh"})
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        assert "Sale-Deed-Ramesh.pdf" in response.headers["content-disposition"]
        assert response.content[:5] == b"%PDF-"

    def test_pdf_export_prints_the_sections_the_user_approved(self, client):
        sections = [dict(s) for s in build_sections(FULL_FIELDS)]
        sections[1]["body"] = "This preamble was edited before download."
        response = client.post("/drafting/export/pdf",
                               json={"fields": FULL_FIELDS, "sections": sections})
        _, text = read_pdf_text(io.BytesIO(response.content))
        assert "edited before download" in text

    def test_a_dangerous_filename_cannot_escape(self, client):
        response = client.post("/drafting/export/pdf",
                               json={"fields": FULL_FIELDS,
                                     "filename": "../../etc/passwd"})
        disposition = response.headers["content-disposition"]
        assert ".." not in disposition and "/" not in disposition

    def test_docx_export_returns_a_downloadable_word_file(self, client):
        response = client.post("/drafting/export/docx", json={"fields": FULL_FIELDS})
        assert response.status_code == 200
        assert "wordprocessingml" in response.headers["content-type"]
        assert zipfile.is_zipfile(io.BytesIO(response.content))
